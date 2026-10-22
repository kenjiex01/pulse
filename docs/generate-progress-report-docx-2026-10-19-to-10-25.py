#!/usr/bin/env python3
"""Generate user-friendly Word weekly progress report: Oct 12-18 progress / Oct 19-25 targets."""

from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

OUT = Path(__file__).with_name(
    "People360-and-Collector-Progress-and-Target-Report-2026-10-19-to-10-25.docx"
)


def set_run_font(run, *, size=11, bold=False, color=None):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    run.font.size = Pt(size)
    run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14 if level == 1 else 12)
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run(text)
    if level == 1:
        set_run_font(run, size=18, bold=True, color=RGBColor(0x1F, 0x4E, 0x79))
    elif level == 2:
        set_run_font(run, size=13, bold=True, color=RGBColor(0x2E, 0x75, 0xB6))
    else:
        set_run_font(run, size=12, bold=True, color=RGBColor(0x2E, 0x75, 0xB6))
    return p


def add_meta(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    set_run_font(run, size=11, bold=True)
    return p


def add_para(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    set_run_font(run, size=11)
    return p


def add_item(doc, title, body):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.15
    title_run = p.add_run(title)
    set_run_font(title_run, size=11, bold=True)
    body_run = p.add_run(" " + body)
    set_run_font(body_run, size=11)
    return p


def add_numbered(doc, n, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(f"{n}. {text}")
    set_run_font(run, size=11)
    return p


def main():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)

    add_heading(doc, "Weekly Progress Report & Targets", 1)
    add_meta(doc, "Progress Period: October 12 - October 18, 2026")
    add_meta(doc, "Targets Period: October 19 - October 25, 2026")
    add_meta(doc, "Apps: People360 (school payroll and timekeeping) and Pulse Biometric Collector")

    add_heading(doc, "In plain words", 2)
    add_para(
        doc,
        "Last week we shipped People360 through version 1.0.36 with a staff user guide "
        "(screenshots and payroll workflow order), payroll batch fixes (add employees by pay-period "
        "category, paging, staff vs faculty hours by category), loan deductions and history "
        "(one deduction per loan per run, optional half per cutoff, clearer principal and interest), "
        "timekeeping fixes (overnight shifts, overtime in whole minutes), payslip year-to-date totals, "
        "and nine ICCT Company Document memos and letters with preview and email PDF matching the screen. "
        "HR lookup lists gained template download and CSV upload. Desktop template downloads save real "
        "Excel files to Downloads. Pulse Collector 1.0.5 fixes DTR export to save CSV instead of hanging.",
    )
    add_para(
        doc,
        "This week campuses should update to People360 1.0.36, run HR tests on memos, payslip YTD, "
        "loans, and lookup uploads, and share the user guide. Continue Pulse 1.0.5 rollout. Approved "
        "leave from People360 web into desktop payroll stays the main build focus when HR is ready.",
    )

    add_heading(doc, "What we finished last week — People360", 2)

    add_heading(doc, "Staff guide and desktop reliability", 3)
    add_item(
        doc,
        "People360 User Guide (Word).",
        "English guide with screenshots: sign-in, CRUD pattern, setup order, employees, timekeeping, payroll, admin; payroll workflow order.",
    )
    add_item(
        doc,
        "Excel templates on desktop.",
        "Upload templates work without XMLWriter; Download Template saves .xlsx or .csv to Downloads on Windows and Mac.",
    )
    add_item(
        doc,
        "Startup and dev.",
        "Loading overlay hides if CSS fails; dev server auto-start fixed under Documents on Mac.",
    )

    add_heading(doc, "Payroll batch, categories, salary uploads", 3)
    add_item(
        doc,
        "Add employees on a batch.",
        "Respects categories on the pay period; 25 per page with Previous/Next; Include all eligible still works.",
    )
    add_item(
        doc,
        "Employee categories.",
        "User Type lookup; payroll can use Logs or Teaching Loads hours by category.",
    )
    add_item(
        doc,
        "Salary template.",
        "Filter by category; name column on template; upload matches by employee number.",
    )

    add_heading(doc, "Loans and payslips", 3)
    add_item(
        doc,
        "Loan payroll.",
        "Payment history on employee loans; one batch per loan unless half per cutoff; principal, interest, payment count; Deductions tab.",
    )
    add_item(
        doc,
        "Payslips.",
        "Year-to-date figures under signature on posted payslips; PDF layout for email; password birth date MMDDYYYY.",
    )

    add_heading(doc, "Timekeeping and employee maintenance", 3)
    add_item(
        doc,
        "Overnight shifts.",
        "Time-out after midnight stays on the same shift day.",
    )
    add_item(
        doc,
        "Overtime filing.",
        "Filed OT in the same minute as a punch accepted; OT duration in whole minutes.",
    )
    add_item(
        doc,
        "Position and designation.",
        "Upload adds names to maintenance so dropdowns stop showing Legacy for valid titles.",
    )

    add_heading(doc, "Company Documents (HR memos)", 3)
    add_item(
        doc,
        "Nine ICCT templates.",
        "Absences, Attendance Pattern, Company Property, PhilHealth Number, Letters 1–5; bordered offense table from linked offense.",
    )
    add_item(
        doc,
        "Designer and email.",
        "Preview and email use designer layout; Table element; app update replaces old templates (revision 7 seeder).",
    )
    add_item(
        doc,
        "HR lookup uploads.",
        "Download template and upload CSV on campuses, designations, positions, ranks, departments, colleges, programs, document types.",
    )

    add_heading(doc, "Desktop releases", 3)
    add_item(
        doc,
        "People360 1.0.25 through 1.0.36.",
        "Paired macOS and Windows installers and automatic updates; latest ships company documents, lookup uploads, loan fixes, payslip YTD.",
    )

    add_heading(doc, "Pulse Biometric Collector — last week", 2)
    add_item(
        doc,
        "Pulse 1.0.5.",
        "Download DTR saves CSV to Downloads and closes the saving overlay (no stuck screen on Windows).",
    )

    add_heading(doc, "What we plan to do this week", 2)

    add_heading(doc, "People360 — rollout and HR testing", 3)
    add_numbered(doc, 1, "Roll out People360 1.0.36 on active campuses.")
    add_numbered(
        doc,
        2,
        "Company Documents: HR previews memos, sends test email, confirms offense table and merge tags.",
    )
    add_numbered(doc, 3, "Payslips: spot-check year-to-date on a posted batch; confirm PDF password with HR.")
    add_numbered(
        doc,
        4,
        "Loans: verify one deduction per loan, payment history, half/cutoff on a sample employee.",
    )
    add_numbered(doc, 5, "HR lookup uploads: template download and one test upload per module.")
    add_numbered(doc, 6, "Payroll batch: category-scoped add employees; paging and Include all.")
    add_numbered(doc, 7, "Distribute People360 User Guide to campus HR/payroll leads.")
    add_numbered(doc, 8, "Fix campus-reported issues from 1.0.30–1.0.36.")

    add_heading(doc, "People360 — approved leave from web", 3)
    add_numbered(
        doc,
        9,
        "Continue build/test: web approved leave into desktop payroll; mark used after post.",
    )
    add_numbered(doc, 10, "HR test path when ready: web approve → desktop batch → post.")

    add_heading(doc, "Pulse Biometric Collector", 3)
    add_numbered(doc, 11, "Upgrade campuses to Pulse 1.0.5.")
    add_numbered(doc, 12, "Confirm DTR download produces CSV in Downloads on Windows and Mac.")
    add_numbered(
        doc,
        13,
        "Keep 1.0.3 retention/storage practices: 2-month default, device storage column.",
    )
    add_numbered(doc, 14, "Schedule installs on campuses still pending.")

    add_para(doc, "Report date: 19 October 2026.")

    doc.save(OUT)
    print(f"Wrote {OUT}")
    print(f"Size: {OUT.stat().st_size} bytes")


if __name__ == "__main__":
    main()
