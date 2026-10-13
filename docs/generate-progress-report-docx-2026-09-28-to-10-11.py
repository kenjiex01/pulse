#!/usr/bin/env python3
"""Generate user-friendly Word weekly progress report: Sep 28-Oct 4 progress / Oct 5-11 targets."""

from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

OUT = Path(__file__).with_name(
    "People360-and-Collector-Progress-and-Target-Report-2026-09-28-to-10-11.docx"
)


def set_run_font(run, *, size=11, bold=False, color=None):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    run.font.size = Pt(size)
    run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14 if level == 1 else 12)
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run(text)
    if level == 1:
        set_run_font(run, size=18, bold=True, color=RGBColor(0x1F, 0x4E, 0x79))
    elif level == 2:
        set_run_font(run, size=13, bold=True, color=RGBColor(0x2E, 0x75, 0xB6))
    else:
        set_run_font(run, size=12, bold=True, color=RGBColor(0x2E, 0x75, 0xB6))
    return p


def add_meta(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    set_run_font(run, size=11, bold=True)
    return p


def add_para(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    set_run_font(run, size=11)
    return p


def add_item(doc, title, body):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.15
    title_run = p.add_run(title)
    set_run_font(title_run, size=11, bold=True)
    body_run = p.add_run(" " + body)
    set_run_font(body_run, size=11)
    return p


def add_numbered(doc, n, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(f"{n}. {text}")
    set_run_font(run, size=11)
    return p


def main():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)

    add_heading(doc, "Weekly Progress Report & Targets", 1)
    add_meta(doc, "Progress Period: September 28 - October 4, 2026")
    add_meta(doc, "Targets Period: October 5 - October 11, 2026")
    add_meta(doc, "Apps: People360 (school payroll and timekeeping) and Pulse Biometric Collector")

    add_heading(doc, "In plain words", 2)
    add_para(
        doc,
        "Last week we moved campuses toward People360 1.0.14 through 1.0.22: automatic memos when "
        "payroll is posted, refreshed HR letter and memo templates, and fewer Windows firewall pop-ups "
        "for same-network database connect. Employee Profile gained Staff/Faculty filter, per-weekday "
        "shift codes, and smoother bulk upload. HR Setup is new: HR email(s), probationary-end "
        "reminders, notification history, and a daily 5 PM email for campuses with no biometric upload "
        "today. Employee Load templates pull from Skolaris class schedules. Faculty memos can use "
        "teaching load without shift setup; break tardiness follows shift break windows when set. "
        "Pulse Collector 1.0.2 fixes the stuck Loading screen.",
    )
    add_para(
        doc,
        "This week campuses should install or auto-update to People360 1.0.22 and Pulse 1.0.2, then test "
        "with HR: posted-payroll memos, probationary emails (real inbox), employee load templates, "
        "faculty memos, and break-late payroll. Approved leave from People360 web into desktop payroll "
        "remains the main build focus. Collector rollout continues beyond Antipolo.",
    )

    add_heading(doc, "What we finished last week — People360", 2)

    add_heading(doc, "Payroll, memos, and time rules", 3)
    add_item(
        doc,
        "Automatic memos on payroll post.",
        "After post, late/undertime/absent counts vs Memo Setup thresholds trigger memo emails; post is not blocked on send failure.",
    )
    add_item(
        doc,
        "Break tardiness vs shift schedule.",
        "When policy and shift Break Out/In exist, lunch tardiness uses the schedule window, not only total break minutes.",
    )
    add_item(
        doc,
        "Faculty memos from teaching load.",
        "Faculty without shift setup can be included via Employee Load schedules; unmarked Skolaris checker rows count as Present.",
    )

    add_heading(doc, "Company documents and HR letters", 3)
    add_item(
        doc,
        "HR letter and memo templates.",
        "Nine ICCT-aligned forms with correct logo; designer float left/right and nature-of-offense merge tag.",
    )
    add_item(
        doc,
        "Web NTE tracking.",
        "Send creates NTE cases with due dates; HR marks received; Skolaris pull ready when API is live.",
    )

    add_heading(doc, "Employee Profile and Employee Load", 3)
    add_item(
        doc,
        "Staff / Faculty filter and weekly shifts.",
        "Category filter on Employee Profile; different shift code per weekday; rest days exempt from shift requirement.",
    )
    add_item(
        doc,
        "Profile bulk upload fixes.",
        "UTF-8 BOM, skip blank rows, weekly shift CSV columns documented.",
    )
    add_item(
        doc,
        "Employee Load from Skolaris schedules.",
        "Monthly template from daily-loads with faster download and default Time In/Out from class schedule.",
    )

    add_heading(doc, "HR Setup and notifications", 3)
    add_item(
        doc,
        "HR Setup screen.",
        "HR emails, probationary reminder day list, customizable subject/body, notification history tab.",
    )
    add_item(
        doc,
        "Probationary end emails.",
        "First app open each day sends HR-only reminders; warns if mail is log-only.",
    )
    add_item(
        doc,
        "Missing biometric upload (5 PM).",
        "Daily email to HR listing campuses with no upload today (Manila time).",
    )

    add_heading(doc, "Desktop and releases", 3)
    add_item(
        doc,
        "Same-network database (Windows).",
        "Firewall prompt on Database Refresh, not every app open; fewer repeat UAC prompts.",
    )
    add_item(
        doc,
        "People360 1.0.14 through 1.0.22.",
        "Paired macOS and Windows installers and auto-update through 1.0.22.",
    )
    add_item(
        doc,
        "Editable employment Effectivity To.",
        "End date on current employment with validation vs Effectivity From.",
    )

    add_heading(doc, "Pulse Biometric Collector — last week", 2)
    add_item(
        doc,
        "Pulse 1.0.2 released.",
        "Fixes stuck Loading screen: background auto-collect, non-blocking updater check, loader reset on navigation.",
    )
    add_para(
        doc,
        "Antipolo remains the primary installed site; Binangonan, Taytay, and other campuses still pending install or upgrade to 1.0.2.",
    )

    add_heading(doc, "What we plan to do this week", 2)

    add_heading(doc, "People360 — rollout and HR testing", 3)
    add_numbered(doc, 1, "Roll out People360 1.0.22 on all active campuses.")
    add_numbered(doc, 2, "Test automatic memos on payroll post with real Memo Setup counts.")
    add_numbered(doc, 3, "Configure HR Setup and confirm probationary emails reach real inboxes (SMTP).")
    add_numbered(doc, 4, "Verify 5 PM missing-upload email vs dashboard status.")
    add_numbered(doc, 5, "Employee Load: full-month template download, edit, upload; faculty payroll/memo check.")
    add_numbered(doc, 6, "Test break tardiness with shift Break Out/In and shifted lunch punches.")
    add_numbered(doc, 7, "Fix campus issues from 1.0.14–1.0.22.")

    add_heading(doc, "People360 — approved leave from web", 3)
    add_numbered(doc, 8, "Import approved filed leave from People360 web.")
    add_numbered(doc, 9, "Apply in desktop payroll batch; mark web forms used after post.")
    add_numbered(doc, 10, "HR test: approve on web → desktop batch → post.")

    add_heading(doc, "People360 — Company documents / NTE", 3)
    add_numbered(doc, 11, "Pilot web NTE with HR; enable Skolaris pull when API is ready.")

    add_heading(doc, "Pulse Biometric Collector", 3)
    add_numbered(doc, 12, "Upgrade Antipolo to Pulse 1.0.2; confirm loader and auto-collect.")
    add_numbered(doc, 13, "Schedule first install on Binangonan, Taytay, and other campuses.")
    add_numbered(doc, 14, "Confirm retention months saved; UI stays responsive during collect.")

    add_para(doc, "Report date: 5 October 2026.")

    doc.save(OUT)
    print(f"Wrote {OUT}")
    print(f"Size: {OUT.stat().st_size} bytes")


if __name__ == "__main__":
    main()
