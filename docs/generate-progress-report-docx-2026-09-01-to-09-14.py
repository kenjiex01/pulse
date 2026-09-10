#!/usr/bin/env python3
"""Generate user-friendly Word weekly progress report: Sep 1-7 progress / Sep 8-14 targets."""

from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

OUT = Path(__file__).with_name(
    "People360-and-Collector-Progress-and-Target-Report-2026-09-01-to-09-14.docx"
)


def set_run_font(run, *, size=11, bold=False, color=None):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    run.font.size = Pt(size)
    run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14 if level == 1 else 12)
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run(text)
    if level == 1:
        set_run_font(run, size=18, bold=True, color=RGBColor(0x1F, 0x4E, 0x79))
    elif level == 2:
        set_run_font(run, size=13, bold=True, color=RGBColor(0x2E, 0x75, 0xB6))
    else:
        set_run_font(run, size=12, bold=True, color=RGBColor(0x2E, 0x75, 0xB6))
    return p


def add_meta(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    set_run_font(run, size=11, bold=True)
    return p


def add_para(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    set_run_font(run, size=11)
    return p


def add_item(doc, title, body):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.15
    title_run = p.add_run(title)
    set_run_font(title_run, size=11, bold=True)
    body_run = p.add_run(" " + body)
    set_run_font(body_run, size=11)
    return p


def add_numbered(doc, n, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(f"{n}. {text}")
    set_run_font(run, size=11)
    return p


def main():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)

    add_heading(doc, "Weekly Progress Report & Targets", 1)
    add_meta(doc, "Progress Period: September 1 - September 7, 2026")
    add_meta(doc, "Targets Period: September 8 - September 14, 2026")
    add_meta(doc, "Apps: People360 (school payroll and timekeeping) and Biometric Collector")

    add_heading(doc, "In plain words", 2)
    add_para(
        doc,
        "Last week we shipped payslip email from Payroll Transaction — HR can select a posted batch "
        "and employees, and each person receives their payslip PDF through AWS SES "
        "(noreply@skolaris.icct.edu.ph).",
    )
    add_para(
        doc,
        "This week the focus is pulling approved forms from People360 web (OT amendments, Job Order / WRF, "
        "and similar ESS requests) into the desktop payroll batch so payroll does not re-key what managers "
        "already approved on web.",
    )

    add_heading(doc, "What we finished last week — People360", 2)

    add_heading(doc, "Payslip email", 3)
    add_item(
        doc,
        "Payroll Transaction → Payslip tab.",
        "After post, HR selects batch and employees, confirms, and sends with a progress bar (e.g. 3 / 10).",
    )
    add_item(
        doc,
        "PDF attachment.",
        "Each employee gets a short message plus payslip PDF; net pay shown in the email body. Requires email on master record.",
    )
    add_item(
        doc,
        "AWS SES.",
        "Uses SES API (not SMTP) with sender noreply@skolaris.icct.edu.ph in ap-southeast-2. Clearer errors when AWS auth fails.",
    )

    add_heading(doc, "Biometric Collector — last week", 2)
    add_para(doc, "No new Collector installer was released this week.")
    add_para(
        doc,
        "Antipolo only (earlier version). Binangonan, Taytay, and other campuses do not have Collector yet.",
    )

    add_heading(doc, "What we plan to do this week", 2)
    add_heading(doc, "People360 — approved web forms → payroll batch", 3)
    add_numbered(
        doc,
        1,
        "Pull approved ESS forms from People360 web via pulse-api employee-requests sync; "
        "local store + Sync from web action (OT amendments, Job Order, WRF, etc.).",
    )
    add_numbered(
        doc,
        2,
        "Apply synced forms to the payroll batch (attendance/OT lines, adjustments); "
        "mark web rows posted after payroll post where applicable.",
    )
    add_numbered(
        doc,
        3,
        "HR UAT on staging with sample approved forms; fix findings and audit each import.",
    )

    add_heading(doc, "Dependencies", 3)
    add_item(doc, "Pulse API key (employee_requests).", "Skolaris admin — needed for web form sync.")
    add_item(doc, "Sample approved OT + Job Order.", "HR — needed for payroll batch UAT.")
    add_item(doc, "Campus PCs 1.0.2+ with HTTPS.", "Campus IT — needed for web sync.")

    add_heading(doc, "Biometric Collector", 3)
    add_numbered(doc, 4, "Antipolo — upgrade to new auto-update installer.")
    add_numbered(doc, 5, "Other campuses — first-time install (Binangonan, Taytay, etc.).")

    doc.save(OUT)
    print(f"Wrote {OUT}")
    print(f"Size: {OUT.stat().st_size} bytes")


if __name__ == "__main__":
    main()
