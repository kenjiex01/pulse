#!/usr/bin/env python3
"""Generate user-friendly Word weekly progress report: Sep 14-20 progress / Sep 21-27 targets."""

from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

OUT = Path(__file__).with_name(
    "People360-and-Collector-Progress-and-Target-Report-2026-09-14-to-09-27.docx"
)


def set_run_font(run, *, size=11, bold=False, color=None):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    run.font.size = Pt(size)
    run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14 if level == 1 else 12)
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run(text)
    if level == 1:
        set_run_font(run, size=18, bold=True, color=RGBColor(0x1F, 0x4E, 0x79))
    elif level == 2:
        set_run_font(run, size=13, bold=True, color=RGBColor(0x2E, 0x75, 0xB6))
    else:
        set_run_font(run, size=12, bold=True, color=RGBColor(0x2E, 0x75, 0xB6))
    return p


def add_meta(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    set_run_font(run, size=11, bold=True)
    return p


def add_para(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    set_run_font(run, size=11)
    return p


def add_item(doc, title, body):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.15
    title_run = p.add_run(title)
    set_run_font(title_run, size=11, bold=True)
    body_run = p.add_run(" " + body)
    set_run_font(body_run, size=11)
    return p


def add_numbered(doc, n, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(f"{n}. {text}")
    set_run_font(run, size=11)
    return p


def main():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)

    add_heading(doc, "Weekly Progress Report & Targets", 1)
    add_meta(doc, "Progress Period: September 14 - September 20, 2026")
    add_meta(doc, "Targets Period: September 21 - September 27, 2026")
    add_meta(doc, "Apps: People360 (school payroll and timekeeping) and Biometric Collector")

    add_heading(doc, "In plain words", 2)
    add_para(
        doc,
        "Last week we continued Company Documents (MEMO) so HR can send a memo from the template "
        "card, attach a fillable Notice to Explain in Word, keep a send history, and download a "
        "Memo report (who received which document, when, and which offense number). Offense "
        "categories and penalties can now be maintained in the app. People360 1.0.7 was released. "
        "Master File upload also became safer: it matches by employee number, does not wipe blank "
        "cells, and no longer fails when an employee has more than one campus.",
    )
    add_para(
        doc,
        "This week campuses should install or auto-update to People360 1.0.7 and test memos. "
        "The main new build work is pulling approved filed leave from People360 web into desktop "
        "payroll, so HR does not re-encode leave that managers already approved. We will also "
        "continue Collector rollout beyond Antipolo.",
    )

    add_heading(doc, "What we finished last week — People360", 2)

    add_heading(doc, "Company Documents (MEMO)", 3)
    add_item(
        doc,
        "Send from the template card.",
        "HR clicks Send on an active memo, picks employees (search and Select all), and emails "
        "the PDF. A progress bar shows how many have been sent (for example 27 / 250).",
    )
    add_item(
        doc,
        "Send history.",
        "The same screen lists who already received that document, how many times, and on which "
        "dates. Re-sending the same person adds a new history row.",
    )
    add_item(
        doc,
        "Notice to Explain (NTE).",
        "A memo can be marked Requires NTE or Set as NTE. Only one active template can be the "
        "NTE itself. When NTE is required, the employee email includes the memo PDF and a Word "
        "(.docx) Notice to Explain that they can fill in.",
    )
    add_item(
        doc,
        "Nature of Offense catalog.",
        "HR can maintain the ICCT Code of Offenses list. Memo templates can pick the offense "
        "from that list instead of typing it freehand.",
    )
    add_item(
        doc,
        "Offense categories and penalties.",
        "Offense Categories follows the official A-D penalty table. Memo tags can fill "
        "Disciplinary action and Offense frequency based on how many times that memo was already "
        "sent to the same employee.",
    )
    add_item(
        doc,
        "Memo report.",
        "Reports → Human Resource → Memo. HR picks a date range and company documents (Select all). "
        "The download lists employee, document, send time, sender, and offense frequency.",
    )
    add_item(
        doc,
        "Designer polish.",
        "Canvas and text colors, full legal-size PDF, empty field labels stay empty, and "
        "Timekeeping Send asks for confirmation.",
    )
    add_item(
        doc,
        "People360 1.0.4 through 1.0.7 released.",
        "The latest paired installers are 1.0.7 (macOS and Windows). Campuses on auto-update "
        "can pick up 1.0.7.",
    )

    add_heading(doc, "Employee records", 3)
    add_item(
        doc,
        "Safer Master File upload.",
        "Upload matches existing people by employee number. Only filled Excel columns change — "
        "blank cells do not erase what is already on file.",
    )
    add_item(
        doc,
        "Upload changes in history.",
        "Master File updates now show on Employee History and in the Historical Data report.",
    )
    add_item(
        doc,
        "Multi-campus employees.",
        "Uploading campuses no longer errors when the employee already has more than one campus.",
    )

    add_heading(doc, "Desktop app", 3)
    add_item(
        doc,
        "One window on Windows.",
        "Opening reports, uploads, or links no longer pops a second People360 window. Outside "
        "websites still open in the normal browser.",
    )

    add_heading(doc, "Biometric Collector — last week", 2)
    add_para(doc, "No new Collector installer was released this week.")
    add_para(
        doc,
        "Antipolo only (earlier version). Binangonan, Taytay, and other campuses do not have Collector yet.",
    )

    add_heading(doc, "What we plan to do this week", 2)

    add_heading(doc, "People360 — campus use of MEMO on 1.0.7", 3)
    add_numbered(
        doc,
        1,
        "Roll out People360 1.0.7 so campuses have Send, NTE Word, send history, offense tags, "
        "and the Memo report.",
    )
    add_numbered(
        doc,
        2,
        "HR test of send and NTE: preview, send a sample, confirm the PDF matches the screen, "
        "and that the Word NTE opens when required.",
    )
    add_numbered(
        doc,
        3,
        "HR test of the Memo report: date range, selected documents, Select all; confirm employees, "
        "documents, send times, and offense number.",
    )
    add_numbered(
        doc,
        4,
        "Fix findings from campus use of 1.0.7.",
    )

    add_heading(doc, "People360 — approved leave from web into payroll", 3)
    add_numbered(
        doc,
        5,
        "Pull approved filed leave from People360 web (employee, leave type, dates, hours/days) "
        "so payroll does not re-type what managers already approved.",
    )
    add_numbered(
        doc,
        6,
        "Process those leaves in the desktop payroll batch. After the batch is posted, mark the "
        "web leave forms as already used in payroll so they are not applied twice.",
    )
    add_numbered(
        doc,
        7,
        "Test with HR: approved leave on web → pull into desktop → appears on the open payroll "
        "batch → post. Fix wrong dates, leave type, or missing employee.",
    )

    add_heading(doc, "Biometric Collector", 3)
    add_numbered(doc, 8, "Antipolo — upgrade to the auto-update installer if not done yet.")
    add_numbered(doc, 9, "Other campuses — first-time install (Binangonan, Taytay, and scheduled sites).")

    doc.save(OUT)
    print(f"Wrote {OUT}")
    print(f"Size: {OUT.stat().st_size} bytes")


if __name__ == "__main__":
    main()
