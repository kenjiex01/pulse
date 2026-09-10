#!/usr/bin/env python3
"""Generate user-friendly Word weekly progress report: Aug 25-31 / Sep 1-7."""

from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

OUT = Path(__file__).with_name(
    "People360-and-Collector-Progress-and-Target-Report-2026-08-25-to-09-07.docx"
)


def set_run_font(run, *, size=11, bold=False, color=None):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    run.font.size = Pt(size)
    run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14 if level == 1 else 12)
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run(text)
    if level == 1:
        set_run_font(run, size=18, bold=True, color=RGBColor(0x1F, 0x4E, 0x79))
    elif level == 2:
        set_run_font(run, size=13, bold=True, color=RGBColor(0x2E, 0x75, 0xB6))
    else:
        set_run_font(run, size=12, bold=True, color=RGBColor(0x2E, 0x75, 0xB6))
    return p


def add_meta(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    set_run_font(run, size=11, bold=True)
    return p


def add_para(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    set_run_font(run, size=11)
    return p


def add_item(doc, title, body):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.15
    title_run = p.add_run(title)
    set_run_font(title_run, size=11, bold=True)
    body_run = p.add_run(" " + body)
    set_run_font(body_run, size=11)
    return p


def add_numbered(doc, n, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(f"{n}. {text}")
    set_run_font(run, size=11)
    return p


def main():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)

    add_heading(doc, "Weekly Progress Report & Targets", 1)
    add_meta(doc, "Progress Period: August 25 - August 31, 2026")
    add_meta(doc, "Targets Period: September 1 - September 7, 2026")
    add_meta(doc, "Apps: People360 (school payroll and timekeeping) and Biometric Collector")

    add_heading(doc, "In plain words", 2)
    add_para(
        doc,
        "Last week we renamed the desktop app to People360, released three installer updates, "
        "made the employee list faster and easier to filter, added a memo layout designer for HR, "
        "fixed slow time-log imports, and prepared planning reports that show which HR modules are already ready.",
    )
    add_para(
        doc,
        "This week the focus is testing what we shipped (especially payslips and payroll reports), "
        "fixing anything HR finds wrong, upgrading Antipolo to the new Biometric Collector with automatic updates, "
        "and installing that new Collector on other campuses for the first time.",
    )

    add_heading(doc, "What we finished last week — People360", 2)

    add_heading(doc, "App name, updates, and installers", 3)
    add_item(
        doc,
        "New name: People360.",
        "The desktop app no longer shows Pulse to users. Existing employee and payroll data is kept when upgrading.",
    )
    add_item(
        doc,
        "Automatic updates.",
        "New versions can download and install in the background. Campus PCs no longer need special setup just to check for updates.",
    )
    add_item(
        doc,
        "Installers released:",
        "Version 1.0.0 (first People360 release), 1.0.1 (employee list improvements), 1.0.2 (easier updates on all PCs).",
    )

    add_heading(doc, "Employee records — easier to find people", 3)
    add_item(
        doc,
        "New filters.",
        "HR can narrow the list by campus, department or college, and faculty vs staff.",
    )
    add_item(
        doc,
        "Much faster list.",
        "The employee screen used to take about half a minute; it now loads quickly when filters change.",
    )
    add_item(
        doc,
        "HR Employee report.",
        "Focuses on personal info, assignments, salary, and loans — without mixed-in document columns.",
    )

    add_heading(doc, "Company Documents — design memo layouts", 3)
    add_item(
        doc,
        "New HR screen.",
        "HR can create memo templates with a visual designer: drag fields, text, images, signatures, and dates.",
    )
    add_item(
        doc,
        "Current scope.",
        "Template design only for now. Sending memos for approval through the app is not turned on yet.",
    )

    add_heading(doc, "Timekeeping", 3)
    add_item(
        doc,
        "Pulling time clock files.",
        "Importing a full month no longer crashes with a timeout. Connection problems show a clear message.",
    )
    add_item(
        doc,
        "Time Logs paging.",
        "Long batch lists now have First / Previous / Next / Last buttons.",
    )

    add_heading(doc, "HR planning documents", 3)
    add_item(
        doc,
        "Module checklist.",
        "Mapped 28 standard HR modules against what People360 desktop and web already have.",
    )
    add_item(
        doc,
        "Three modules ready on desktop.",
        "Organization setup, Time and Attendance processing, and Leave for payroll. Employee leave filing and mobile clock-in are planned for web.",
    )

    add_heading(doc, "Biometric Collector — last week", 2)
    add_para(doc, "No new Collector installer was released this week.")
    add_para(
        doc,
        "Installed so far: Biometric Collector is on Antipolo only, using an earlier version — not yet the new installer with automatic updates. "
        "Binangonan, Taytay, and other campuses do not have Collector installed yet.",
    )

    add_heading(doc, "What we plan to do this week", 2)
    add_heading(doc, "People360", 3)
    add_numbered(
        doc,
        1,
        "Test the latest app (1.0.2+) on campus PCs — employee filters, memo designer, time log paging. "
        "Release 1.0.3 only if something important is broken.",
    )
    add_numbered(
        doc,
        2,
        "Payslip review — run payslips from posted payroll for staff, faculty, and hybrid employees "
        "(on screen, Excel, PDF). Fix wrong layouts, rates, or deductions from HR feedback.",
    )
    add_numbered(
        doc,
        3,
        "Other payroll reports — apply fixes from HR testing on Payroll Register and government reports "
        "(SSS, PhilHealth, Pag-IBIG, BIR).",
    )

    add_heading(doc, "Biometric Collector", 3)
    add_numbered(
        doc,
        4,
        "Antipolo — replace the current Collector with the latest installer that has automatic updates. "
        "Antipolo is the only campus that has Collector installed today.",
    )
    add_numbered(
        doc,
        5,
        "Other campuses — first-time install of the same new auto-update Collector on Binangonan, Taytay, "
        "and other scheduled sites. These campuses do not have Collector yet.",
    )

    doc.save(OUT)
    print(f"Wrote {OUT}")
    print(f"Size: {OUT.stat().st_size} bytes")


if __name__ == "__main__":
    main()
