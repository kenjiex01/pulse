#!/usr/bin/env python3
"""Generate user-friendly Word weekly progress report: Oct 5-11 progress / Oct 12-18 targets."""

from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

OUT = Path(__file__).with_name(
    "People360-and-Collector-Progress-and-Target-Report-2026-10-12-to-10-18.docx"
)


def set_run_font(run, *, size=11, bold=False, color=None):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    run.font.size = Pt(size)
    run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14 if level == 1 else 12)
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run(text)
    if level == 1:
        set_run_font(run, size=18, bold=True, color=RGBColor(0x1F, 0x4E, 0x79))
    elif level == 2:
        set_run_font(run, size=13, bold=True, color=RGBColor(0x2E, 0x75, 0xB6))
    else:
        set_run_font(run, size=12, bold=True, color=RGBColor(0x2E, 0x75, 0xB6))
    return p


def add_meta(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    set_run_font(run, size=11, bold=True)
    return p


def add_para(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    set_run_font(run, size=11)
    return p


def add_item(doc, title, body):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.15
    title_run = p.add_run(title)
    set_run_font(title_run, size=11, bold=True)
    body_run = p.add_run(" " + body)
    set_run_font(body_run, size=11)
    return p


def add_numbered(doc, n, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(f"{n}. {text}")
    set_run_font(run, size=11)
    return p


def main():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)

    add_heading(doc, "Weekly Progress Report & Targets", 1)
    add_meta(doc, "Progress Period: October 5 - October 11, 2026")
    add_meta(doc, "Targets Period: October 12 - October 18, 2026")
    add_meta(doc, "Apps: People360 (school payroll and timekeeping) and Pulse Biometric Collector")

    add_heading(doc, "In plain words", 2)
    add_para(
        doc,
        "Last week we shipped People360 1.0.23 and 1.0.24 with stronger faculty teaching-load "
        "payroll (faster Skolaris pull, every selected faculty listed including zero rows, Present "
        "days pay without punches, red batch names follow teaching-load days). PhilHealth and other "
        "statutory deductions use month-to-date pay through the cutoff being processed. Overtime "
        "policy now separates regular vs non-regular days and respects require filing before auto OT. "
        "Desktop database restore is safer and accepts larger SQL backups. Pulse Collector 1.0.3 "
        "defaults log retention to 2 months and shows device attendance storage when online.",
    )
    add_para(
        doc,
        "This week campuses should update to People360 1.0.24 and Pulse 1.0.3, then re-test payroll "
        "with HR: faculty batches, PhilHealth on split cutoffs, and staff OT when filing is required. "
        "Collector sites should confirm retention and watch device storage. Approved leave from "
        "People360 web into desktop payroll remains a priority build item.",
    )

    add_heading(doc, "What we finished last week — People360", 2)

    add_heading(doc, "Faculty teaching load and payroll batch", 3)
    add_item(
        doc,
        "Teaching load pull from Skolaris.",
        "Faster bulk pull, cancel long runs, batch lists all selected faculty (zero rows in red), search in batch modal.",
    )
    add_item(
        doc,
        "Faculty pay when marked Present.",
        "Schedule days marked Present can pay without time in/out on that date.",
    )
    add_item(
        doc,
        "Red names on posted batch.",
        "Faculty expected days follow teaching-load dates in the period; no-load zero-basic stays red.",
    )
    add_item(
        doc,
        "People360 1.0.23 and 1.0.24.",
        "Paired macOS and Windows installers and automatic updates through 1.0.24.",
    )

    add_heading(doc, "Statutory deductions and overtime policy", 3)
    add_item(
        doc,
        "PhilHealth / SSS / Pag-IBIG month-to-date.",
        "Earlier cutoff in a month no longer uses basic from a later unposted batch; second cutoff can true up.",
    )
    add_item(
        doc,
        "Overtime policy — regular vs non-regular days.",
        "Separate settings; payroll skips auto OT from excess when filing is required for that day type.",
    )

    add_heading(doc, "Database and dev reliability", 3)
    add_item(
        doc,
        "Safer SQL restore.",
        "Failed desktop import rolls back instead of empty database; duplicate roles renamed on import.",
    )
    add_item(
        doc,
        "Larger SQL backup uploads.",
        "Restore uploads up to about 512 MB where server limits allow.",
    )

    add_heading(doc, "Pulse Biometric Collector — last week", 2)
    add_item(
        doc,
        "Pulse 1.0.3 released.",
        "Default 2-month log retention; dashboard shows attendance used / capacity / free when device is online.",
    )

    add_heading(doc, "What we plan to do this week", 2)

    add_heading(doc, "People360 — rollout and HR testing", 3)
    add_numbered(doc, 1, "Roll out People360 1.0.24 on all active campuses.")
    add_numbered(doc, 2, "Reprocess or spot-check PhilHealth on first cutoff vs same-month second batch.")
    add_numbered(doc, 3, "Faculty payroll QA: pull batch, red/black names, Present without punches.")
    add_numbered(doc, 4, "Staff OT: set For Regular Days policy, reprocess sample; confirm filing vs auto excess.")
    add_numbered(doc, 5, "Fix campus issues from 1.0.23–1.0.24.")
    add_numbered(doc, 6, "Optional: test large SQL restore and rollback in a test environment.")

    add_heading(doc, "People360 — approved leave from web", 3)
    add_numbered(doc, 7, "Continue build/test: web approved leave into desktop payroll; mark used after post.")
    add_numbered(doc, 8, "HR test path when ready: web approve → desktop batch → post.")

    add_heading(doc, "Pulse Biometric Collector", 3)
    add_numbered(doc, 9, "Upgrade campuses to Pulse 1.0.3.")
    add_numbered(doc, 10, "Confirm log retention (2 months default) and prune after collect.")
    add_numbered(doc, 11, "Use device storage column before terminals fill; note collector does not clear device memory.")
    add_numbered(doc, 12, "Schedule first install on campuses still pending.")

    add_para(doc, "Report date: 12 October 2026.")

    doc.save(OUT)
    print(f"Wrote {OUT}")
    print(f"Size: {OUT.stat().st_size} bytes")


if __name__ == "__main__":
    main()
