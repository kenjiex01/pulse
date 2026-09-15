#!/usr/bin/env python3
"""Generate user-friendly Word weekly progress report: Sep 8-13 progress / Sep 14-20 targets."""

from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

OUT = Path(__file__).with_name(
    "People360-and-Collector-Progress-and-Target-Report-2026-09-08-to-09-20.docx"
)


def set_run_font(run, *, size=11, bold=False, color=None):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    run.font.size = Pt(size)
    run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14 if level == 1 else 12)
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run(text)
    if level == 1:
        set_run_font(run, size=18, bold=True, color=RGBColor(0x1F, 0x4E, 0x79))
    elif level == 2:
        set_run_font(run, size=13, bold=True, color=RGBColor(0x2E, 0x75, 0xB6))
    else:
        set_run_font(run, size=12, bold=True, color=RGBColor(0x2E, 0x75, 0xB6))
    return p


def add_meta(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    set_run_font(run, size=11, bold=True)
    return p


def add_para(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    set_run_font(run, size=11)
    return p


def add_item(doc, title, body):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.15
    title_run = p.add_run(title)
    set_run_font(title_run, size=11, bold=True)
    body_run = p.add_run(" " + body)
    set_run_font(body_run, size=11)
    return p


def add_numbered(doc, n, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(f"{n}. {text}")
    set_run_font(run, size=11)
    return p


def main():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)

    add_heading(doc, "Weekly Progress Report & Targets", 1)
    add_meta(doc, "Progress Period: September 8 - September 13, 2026")
    add_meta(doc, "Targets Period: September 14 - September 20, 2026")
    add_meta(doc, "Apps: People360 (school payroll and timekeeping) and Biometric Collector")

    add_heading(doc, "In plain words", 2)
    add_para(
        doc,
        "Last week we shipped Company Documents (MEMO) so HR can preview a memo on legal paper "
        "and email the same PDF the employee receives. Ready-made ICCT disciplinary templates "
        "are in the app, and People360 1.0.3 was released. Teaching Loads can now pull uploaded "
        "faculty schedules from Skolaris, and employee Sex is stored as Male / Female.",
    )
    add_para(
        doc,
        "This week we continue Company Documents (MEMO) editing and additional features — "
        "refine templates in the designer, add the next memo capabilities HR needs, and support "
        "campus testing of the memos already released.",
    )

    add_heading(doc, "What we finished last week — People360", 2)

    add_heading(doc, "Company Documents (MEMO)", 3)
    add_item(
        doc,
        "Preview matches the emailed memo.",
        "HR sees the same legal-size page (8.5 x 14 in) that is attached to the employee email. "
        "Field boxes, stamps, and margins stay aligned.",
    )
    add_item(
        doc,
        "Faster, more reliable send.",
        "Sending no longer times out when the template has stamps or images. Batch send can finish "
        "without the app cutting off at 30 seconds.",
    )
    add_item(
        doc,
        "ICCT Code of Offenses templates.",
        "Eight ready templates: Notice to Explain, Verbal Reprimand, Written Warning, Notice of "
        "Suspension, Notice of Dismissal, Return to Work, AWOL notice, and Uniform / ID / Nameplate.",
    )
    add_item(
        doc,
        "Table of Penalties.",
        "The official ICCT penalty matrix (Categories A-D by frequency) is stored so those templates stay consistent.",
    )
    add_item(
        doc,
        "People360 1.0.3 released.",
        "macOS and Windows installers include memo preview/email and earlier payslip email. "
        "Campuses on auto-update can pick up 1.0.3.",
    )

    add_heading(doc, "Teaching Loads", 3)
    add_item(
        doc,
        "Uploaded faculty loading from Skolaris.",
        "Time Logs → Teaching Loads has an Uploaded PDFs tab. People360 pulls faculty loading "
        "already uploaded in Skolaris.",
    )
    add_item(
        doc,
        "Attendance Checker marks on pull.",
        "Skolaris teaching-load pull now includes attendance marks already recorded in Attendance Checker.",
    )

    add_heading(doc, "Employee records", 3)
    add_item(
        doc,
        "Sex saved as Male / Female.",
        "Employee Profile and Master File upload store Male and Female. Existing records were updated. "
        "Upload still accepts lowercase and converts it.",
    )

    add_heading(doc, "Biometric Collector — last week", 2)
    add_para(doc, "No new Collector installer was released this week.")
    add_para(
        doc,
        "Antipolo only (earlier version). Binangonan, Taytay, and other campuses do not have Collector yet.",
    )

    add_heading(doc, "What we plan to do this week", 2)
    add_heading(doc, "People360 — continue Company Documents (MEMO)", 3)
    add_numbered(
        doc,
        1,
        "Keep editing MEMO templates in the Company Documents designer (layout, wording, stamps, "
        "signatures, and merge tags) so templates match how HR writes memos on paper.",
    )
    add_numbered(
        doc,
        2,
        "Additional MEMO features after last week's send/preview release — more template options, "
        "easier editing, and send/preview fixes found in campus use.",
    )
    add_numbered(
        doc,
        3,
        "Campus test of released memos on People360 1.0.3: preview, send a sample, and confirm "
        "the employee PDF matches the screen. Fix anything that looks wrong.",
    )

    add_heading(doc, "Biometric Collector", 3)
    add_numbered(doc, 4, "Antipolo — upgrade to the auto-update installer if not done yet.")
    add_numbered(doc, 5, "Other campuses — first-time install (Binangonan, Taytay, and scheduled sites).")

    doc.save(OUT)
    print(f"Wrote {OUT}")
    print(f"Size: {OUT.stat().st_size} bytes")


if __name__ == "__main__":
    main()
