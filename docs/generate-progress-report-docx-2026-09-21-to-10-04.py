#!/usr/bin/env python3
"""Generate user-friendly Word weekly progress report: Sep 21-27 progress / Sep 28-Oct 4 targets."""

from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

OUT = Path(__file__).with_name(
    "People360-and-Collector-Progress-and-Target-Report-2026-09-21-to-10-04.docx"
)


def set_run_font(run, *, size=11, bold=False, color=None):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    run.font.size = Pt(size)
    run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14 if level == 1 else 12)
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run(text)
    if level == 1:
        set_run_font(run, size=18, bold=True, color=RGBColor(0x1F, 0x4E, 0x79))
    elif level == 2:
        set_run_font(run, size=13, bold=True, color=RGBColor(0x2E, 0x75, 0xB6))
    else:
        set_run_font(run, size=12, bold=True, color=RGBColor(0x2E, 0x75, 0xB6))
    return p


def add_meta(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    set_run_font(run, size=11, bold=True)
    return p


def add_para(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    set_run_font(run, size=11)
    return p


def add_item(doc, title, body):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.15
    title_run = p.add_run(title)
    set_run_font(title_run, size=11, bold=True)
    body_run = p.add_run(" " + body)
    set_run_font(body_run, size=11)
    return p


def add_numbered(doc, n, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(f"{n}. {text}")
    set_run_font(run, size=11)
    return p


def main():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)

    add_heading(doc, "Weekly Progress Report & Targets", 1)
    add_meta(doc, "Progress Period: September 21 - September 27, 2026")
    add_meta(doc, "Targets Period: September 28 - October 4, 2026")
    add_meta(doc, "Apps: People360 (school payroll and timekeeping) and Biometric Collector")

    add_heading(doc, "In plain words", 2)
    add_para(
        doc,
        "Last week we shipped payroll and employee-record improvements: holiday pay on time-log "
        "payroll, a Fixed Rate option so some employees keep full basic pay without "
        "late/undertime/absent cuts, and clearer salary and employment history with effectivity "
        "dates. The dashboard opens faster with skeleton placeholders while data loads. People360 "
        "1.0.8 through 1.0.13 went out with same-network discovery so one computer can list other "
        "People360 desktops and connect to another campus database. Time Logs can auto-pull new "
        "biometric files from cloud backup while the app is open. Memo Setup now has a Count per "
        "violation type so HR sets how many occurrences are needed before a memo is sent.",
    )
    add_para(
        doc,
        "This week campuses should move to People360 1.0.13 and test automatic memos when payroll "
        "is posted (using Memo Setup counts), plus holiday pay and fixed-rate payroll on real "
        "batches. The main build carry-over is approved filed leave from People360 web into desktop "
        "payroll. Collector rollout beyond Antipolo continues.",
    )

    add_heading(doc, "What we finished last week — People360", 2)

    add_heading(doc, "Payroll and holidays", 3)
    add_item(
        doc,
        "Holiday pay (HOLI income).",
        "Holiday income on legal or special holidays when rules apply—including unworked holiday "
        "pay with adjacent time logs and premiums when the employee worked the holiday.",
    )
    add_item(
        doc,
        "Fixed Rate salary.",
        "HR marks Fixed Rate to keep configured Basic Income and skip late, undertime, and absent "
        "reductions. Government deductions still post when the batch includes them.",
    )
    add_item(
        doc,
        "Last Payroll Date.",
        "Payroll computation stops at Last Payroll Date even if the batch period is longer.",
    )

    add_heading(doc, "Employee records", 3)
    add_item(
        doc,
        "Salary history effectivity.",
        "A new Effectivity From closes the previous salary row the day before the new date.",
    )
    add_item(
        doc,
        "Employment history.",
        "Changes archive to Previous Employment with effectivity dates. Separation Date follows "
        "the last time log and cannot exceed Last Payroll Date.",
    )

    add_heading(doc, "Dashboard and desktop", 3)
    add_item(
        doc,
        "Faster dashboard.",
        "The home screen loads first; Biometric Collector / cloud backup status loads afterward.",
    )
    add_item(
        doc,
        "Skeleton loading.",
        "Tables and status panels show skeleton placeholders while data loads.",
    )
    add_item(
        doc,
        "People360 1.0.8 through 1.0.13 released.",
        "Paired macOS and Windows installers and auto-update through 1.0.13.",
    )

    add_heading(doc, "Same-network People360", 3)
    add_item(
        doc,
        "Find other People360 computers.",
        "Database lists other open People360 desktops on the same office network.",
    )
    add_item(
        doc,
        "Connect to another computer's database.",
        "An admin can connect, work against that data, disconnect, and save back—login stays local.",
    )

    add_heading(doc, "Timekeeping and memos", 3)
    add_item(
        doc,
        "Auto-pull biometric logs from cloud backup.",
        "Optional background import while the app is open when enabled on Time Logs.",
    )
    add_item(
        doc,
        "Memo Setup — Count.",
        "Late, Undertime, and Absent each have a Count threshold before memos are sent manually.",
    )
    add_item(
        doc,
        "Send Document — extra emails.",
        "Additional To addresses when sending a company document.",
    )

    add_heading(doc, "Biometric Collector — last week", 2)
    add_para(doc, "No new Collector installer was released this week.")
    add_para(
        doc,
        "Antipolo only (earlier version). Binangonan, Taytay, and other campuses do not have "
        "Collector yet. Desktop auto-pull from cloud backup reduces manual import where backup is configured.",
    )

    add_heading(doc, "What we plan to do this week", 2)

    add_heading(doc, "People360 — rollout and HR testing", 3)
    add_numbered(doc, 1, "Roll out People360 1.0.13 (holiday pay, Fixed Rate, LAN connect, auto-pull, Memo counts).")
    add_numbered(
        doc,
        2,
        "Test automatic memos on payroll post: set Memo Setup counts, post a batch, confirm emails and attachments.",
    )
    add_numbered(doc, 3, "Test holiday pay on a real batch with legal or special holidays.")
    add_numbered(doc, 4, "Test Fixed Rate employees—no late/undertime/absent reductions.")
    add_numbered(doc, 5, "Fix findings from campus use of 1.0.8–1.0.13.")

    add_heading(doc, "People360 — approved leave from web into payroll", 3)
    add_numbered(
        doc,
        6,
        "Pull approved filed leave from People360 web (employee, type, dates, hours/days).",
    )
    add_numbered(
        doc,
        7,
        "Process leaves in the desktop payroll batch; after post, mark web forms as used in payroll.",
    )
    add_numbered(
        doc,
        8,
        "Test with HR: approved leave on web → desktop batch → post; fix mismatches.",
    )

    add_heading(doc, "People360 — same-network database (IT / central HR)", 3)
    add_numbered(
        doc,
        9,
        "Pilot connect between two PCs on the same LAN; document firewall or Wi‑Fi issues.",
    )

    add_heading(doc, "Biometric Collector", 3)
    add_numbered(doc, 10, "Antipolo — upgrade to the current auto-update path if not done yet.")
    add_numbered(doc, 11, "Other campuses — first-time install (Binangonan, Taytay, scheduled sites).")

    doc.save(OUT)
    print(f"Wrote {OUT}")
    print(f"Size: {OUT.stat().st_size} bytes")


if __name__ == "__main__":
    main()
